"""
Word report generation.

This module creates the export report document.
"""

import os
from tempfile import NamedTemporaryFile

import pandas as pd
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm



from services.plot_service import plot_site_results


def safe_fmt(value, fmt="{:.6f}", na="N/A"):
    try:
        if value is None or pd.isna(value):
            return na
        return fmt.format(value)
    except Exception:
        return na

def build_boxplot_figure(contact_sod):
    fig, ax = plt.subplots(figsize=(6, 4))

    if contact_sod:
        ax.boxplot(
            contact_sod,
            vert=True,
            patch_artist=True,
            boxprops=dict(facecolor="#D9EAF7", color="black"),
            medianprops=dict(color="red", linewidth=2),
            whiskerprops=dict(color="black"),
            capprops=dict(color="black"),
            flierprops=dict(marker="o", markerfacecolor="gray", markersize=5, alpha=0.6),
        )
        ax.set_ylabel("SOD (g/m²/day)")
        ax.set_title("Contact Chamber SOD Distribution")
        ax.set_xticks([1])
        ax.set_xticklabels(["Contact"])
        ax.grid(True, axis="y", alpha=0.25)
    else:
        ax.text(0.5, 0.5, "No contact chamber data", ha="center", va="center")
        ax.set_axis_off()

    fig.tight_layout()
    return fig

def set_cell_text(cell, text):
    cell.text = "" if text is None else str(text)


def shade_and_style_header(row_cells, fill):
    for cell in row_cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)

        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)


def _add_page_number_field(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def _add_total_pages_field(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " NUMPAGES "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def add_header_and_footer(doc, project_number, project_name):
    section = doc.sections[0]

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_p.text = f"Project Number: {project_number} | Project Name: {project_name}"
    for run in header_p.runs:
        run.font.size = Pt(9)
        run.bold = True

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run("Page ")
    run.font.size = Pt(9)
    _add_page_number_field(footer_p)
    run = footer_p.add_run(" of ")
    run.font.size = Pt(9)
    _add_total_pages_field(footer_p)


def group_results_by_type(results):
    grouped = {
        "contact chamber": [],
        "blank chamber": [],
        "ambient": [],
        "other": [],
    }
    for r in results:
        ctype = (r.record.chamber_type or "").strip().lower()
        if ctype in grouped:
            grouped[ctype].append(r)
        else:
            grouped["other"].append(r)
    return grouped


def set_default_font(doc, font_name="Calibri", font_size=12):
    styles = doc.styles

    normal_style = styles["Normal"]
    normal_style.font.name = font_name
    normal_style.font.size = Pt(font_size)

    # Make sure Word uses the font for all script types
    for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        if style_name in styles:
            style = styles[style_name]
            style.font.name = font_name
            style.font.size = Pt(font_size)

            if style.element.rPr is None:
                continue
            rFonts = style.element.rPr.rFonts
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                style.element.rPr.append(rFonts)

            rFonts.set(qn("w:ascii"), font_name)
            rFonts.set(qn("w:hAnsi"), font_name)
            rFonts.set(qn("w:eastAsia"), font_name)
            rFonts.set(qn("w:cs"), font_name)

def build_word_report(path, project_number, project_name, results, blank_mean, site_summary, theta, v_over_a):
    """
    Build the Word document export.

    This preserves the current report structure:
    - header/footer
    - cover page
    - interpretation section
    - grouped result tables
    - plot image
    """
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    #Default font to Calibri
    set_default_font(doc, "Calibri", 12)
    
    add_header_and_footer(doc, project_number, project_name)

    title = doc.add_heading("SOD Calculation Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text, bold in [
        (f"Project Number: {project_number}\n", True),
        (f"Project Name: {project_name}\n", True),
        (f"Report Generated: {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M:%S')}\n", False),
        (f"Files Processed: {len(results)}\n", False),        
        (f"Theta: {theta}\n", False),
        (f"V/A (L/m²): {v_over_a}\n", False),
        (f"Blank Mean Slope: {safe_fmt(blank_mean)} mg/L/min\n", False),
        (f"Mean Site SOD: {safe_fmt(site_summary.get('mean_sod_g_m2_day'))} g/m²/day\n", False),
        (f"Mean SOD corrected to 20°C: {safe_fmt(site_summary.get('mean_sod_20c_g_m2_day'))} g/m²/day", False),
        
    ]:
        run = p.add_run(text)
        run.bold = bold

    doc.add_page_break()

    doc.add_heading("Report Interpretation", level=1)
    doc.add_paragraph(
        "This report summarizes the selected regression windows, calculated slopes, "
        "blank adjustment, temperature correction, and final SOD values for each chamber file."
    )
    doc.add_paragraph(
        "Contact chamber results are used to calculate site-level SOD statistics. "
        "Blank chamber results are used to compute the blank mean slope. "
        "Ambient results are listed separately and are not used in the site mean calculations."
    )
    doc.add_paragraph(
        "The SOD value is derived from the temperature-corrected rate after blank adjustment "
        "and conversion using the selected V/A value."
    )

    doc.add_page_break()

    grouped = group_results_by_type(results)

    doc.add_heading("Contact Chamber Results", level=1)
    contact_results = grouped["contact chamber"]
    if contact_results:
        table = doc.add_table(rows=1, cols=7)
        table.autofit = True
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        headers = ["Site", "Chamber ID", "n", "Slope", "R²", "Adj.DO Rate", "SOD(g/m²/day)"]
        for i, h in enumerate(headers):
            hdr[i].text = h
        shade_and_style_header(hdr, "D9EAF7")

        for r in contact_results:
            row = table.add_row().cells
            values = [
                r.record.site_name,
                r.record.chamber_id,
                r.n_points,
                safe_fmt(r.slope),
                safe_fmt(r.r2),
                safe_fmt(r.adjusted_do_rate),
                safe_fmt(r.sod_g_m2_day),
            ]
            for i, value in enumerate(values):
                set_cell_text(row[i], value)
    else:
        doc.add_paragraph("No contact chamber results were computed.")

    
    doc.add_heading("Contact Chamber SOD Statistics", level=1)
    contact_sod = [r.sod_g_m2_day for r in contact_results if pd.notna(r.sod_g_m2_day)]

    stats_table = doc.add_table(rows=1, cols=2)
    stats_table.autofit = True
    stats_table.style = "Table Grid"
    hdr = stats_table.rows[0].cells
    hdr[0].text = "Statistic"
    hdr[1].text = "SOD (g/m²/day)"
    shade_and_style_header(hdr, "E2F0D9")

    mean_sod = float(pd.Series(contact_sod).mean()) if contact_sod else float("nan")
    min_sod = float(pd.Series(contact_sod).min()) if contact_sod else float("nan")
    max_sod = float(pd.Series(contact_sod).max()) if contact_sod else float("nan")
    median_sod = float(pd.Series(contact_sod).median()) if contact_sod else float("nan")

    for label, value in [
        ("Mean", safe_fmt(mean_sod)),
        ("Median", safe_fmt(median_sod)),
        ("Min", safe_fmt(min_sod)),
        ("Max", safe_fmt(max_sod)),
    ]:
        row = stats_table.add_row().cells
        set_cell_text(row[0], label)
        set_cell_text(row[1], value)

    doc.add_paragraph("")

    fig_box = build_boxplot_figure(contact_sod)

    with NamedTemporaryFile(suffix=".png", delete=False) as tmp_box:
        fig_box.savefig(tmp_box.name, dpi=300, bbox_inches="tight")
        box_path = tmp_box.name

    doc.add_picture(box_path, width=Inches(5.5))
    plt.close(fig_box)

    try:
        os.remove(box_path)
    except OSError:
        pass

    doc.add_heading("Blank Chamber Results", level=1)
    blank_results = grouped["blank chamber"]
    if blank_results:
        table = doc.add_table(rows=1, cols=6)
        table.autofit = True
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        headers = ["Site", "Chamber ID", "N", "Slope", "R²", "Raw Rate"]
        for i, h in enumerate(headers):
            hdr[i].text = h
        shade_and_style_header(hdr, "FCE4D6")

        for r in blank_results:
            row = table.add_row().cells
            values = [
                r.record.site_name,
                r.record.chamber_id,
                r.n_points,
                safe_fmt(r.slope),
                safe_fmt(r.r2),
                safe_fmt(r.unadjusted_do_rate),
            ]
            for i, value in enumerate(values):
                set_cell_text(row[i], value)
    else:
        doc.add_paragraph("No blank chamber results were computed.")

    doc.add_paragraph(f"Blank Mean Slope: {safe_fmt(blank_mean)} mg/L/min")
    doc.add_page_break()

    doc.add_heading("Ambient Results", level=1)
    ambient_results = grouped["ambient"]
    if ambient_results:
        table = doc.add_table(rows=1, cols=5)
        table.autofit = True
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        headers = ["Site", "Chamber ID", "n", "Slope", "R²"]
        for i, h in enumerate(headers):
            hdr[i].text = h
        shade_and_style_header(hdr, "EDEDED")

        for r in ambient_results:
            row = table.add_row().cells
            values = [
                r.record.site_name,
                r.record.chamber_id,
                r.n_points,
                safe_fmt(r.slope),
                safe_fmt(r.r2),
            ]
            for i, value in enumerate(values):
                set_cell_text(row[i], value)
    else:
        doc.add_paragraph("No ambient results were computed.")

    doc.add_paragraph(
        "Ambient records are shown for completeness. They are excluded from blank adjustment and site summary calculations."
    )
    doc.add_page_break()

    doc.add_heading("Regression Plot", level=1)
    fig = plot_site_results(
        results,
        project_number,
        project_name,
        site_summary=site_summary,
        blank_mean=blank_mean,
    )

    with NamedTemporaryFile(suffix=".png", delete=False) as tmp:
        fig.savefig(tmp.name, dpi=300, bbox_inches="tight")
        tmp_path = tmp.name

    doc.add_picture(tmp_path, width=Inches(6.5))
    plt.close(fig)

    doc.save(path)

    try:
        os.remove(tmp_path)
    except OSError:
        pass